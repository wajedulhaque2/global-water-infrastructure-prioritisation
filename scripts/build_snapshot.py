from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Global_Water_Infrastructure_Prioritisation_Tool.xlsx"
OUTPUT = ROOT / "dashboard" / "data"


def extract_table(ws, header_row: int, min_col: int, max_col: int, max_row: int) -> pd.DataFrame:
    headers = [ws.cell(header_row, col).value for col in range(min_col, max_col + 1)]
    rows = []
    for row in ws.iter_rows(
        min_row=header_row + 1,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        if row[0] is None:
            continue
        rows.append(row)
    return pd.DataFrame(rows, columns=headers)


def main() -> None:
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    countries = extract_table(workbook["Country Master"], 5, 1, 12, 222)
    history = extract_table(workbook["PQ Data"], 5, 1, 14, 12565)
    latest = extract_table(workbook["PQ Data"], 5, 19, 32, 1773)
    indicators = extract_table(workbook["Indicator Map"], 5, 1, 13, 14)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    countries.to_csv(OUTPUT / "countries.csv", index=False)
    history.to_csv(OUTPUT / "indicator_history.csv", index=False)
    latest.to_csv(OUTPUT / "latest_indicators.csv", index=False)
    indicators.to_csv(OUTPUT / "indicator_map.csv", index=False)

    retrieved = pd.to_datetime(history["RetrievedAt"], errors="coerce").max()
    metadata = {
        "source": "World Bank World Development Indicators",
        "source_url": "https://data.worldbank.org/indicator",
        "retrieved_at": retrieved.isoformat() if pd.notna(retrieved) else None,
        "analysis_start_year": int(pd.to_numeric(history["Year"]).min()),
        "analysis_end_year": int(pd.to_numeric(history["Year"]).max()),
        "country_count": int(countries["ISO3"].nunique()),
        "indicator_count": int(latest["Indicator ID"].nunique()),
        "history_rows": len(history),
        "latest_rows": len(latest),
    }
    (OUTPUT / "metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
